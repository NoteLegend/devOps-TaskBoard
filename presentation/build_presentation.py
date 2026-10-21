#!/usr/bin/env python3
"""Build the editable capstone deck and its matching PDF from one layout.

Requires python-pptx, reportlab and Pillow. Run from any working directory.
All evidence images remain unmodified; the deck embeds the repository originals.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'presentation'
REPO = 'https://github.com/NoteLegend/devOps-TaskBoard'
DEMO_VIDEO = 'https://drive.google.com/file/d/1ZSxd5RlMJAqJkvLA3E_098BrKF4OgJLT/view?usp=sharing'
RUN = REPO + '/actions/runs/38955563228'
SECURITY_RUN = REPO + '/actions/runs/39583131700'
W, H = 960, 540
INK, MUTED, PAPER = '15243A', '596C80', 'F4F7FA'
BLUE, TEAL, GOLD = '2463EB', '087F76', 'AD650B'
DARK, DARK2, WHITE = '101E32', '1B304B', 'FFFFFF'
PALE, LINE = 'E8EFF8', 'DCE4ED'
FONT, MONO = 'Noto Sans', 'DejaVu Sans Mono'

pdfmetrics.registerFont(TTFont('Noto', '/usr/share/fonts/noto/NotoSans-Regular.ttf'))
pdfmetrics.registerFont(TTFont('NotoBold', '/usr/share/fonts/noto/NotoSans-Bold.ttf'))
pdfmetrics.registerFont(TTFont('Mono', '/usr/share/fonts/TTF/DejaVuSansMono.ttf'))
pdfmetrics.registerFont(TTFont('Symbols', '/usr/share/fonts/TTF/DejaVuSans.ttf'))

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = 'Release Helpdesk | DevOps Capstone'
prs.core_properties.subject = 'Application delivery, security, GitOps and monitored Kubernetes'
prs.core_properties.author = 'Release Helpdesk'
prs.core_properties.keywords = 'DevOps, AWS EKS, Terraform, GitHub Actions, Helm, capstone'
pdf = canvas.Canvas(str(OUT / 'Release-Helpdesk-Presentation.pdf'), pagesize=(W, H))
pdf.setTitle(prs.core_properties.title)
pdf.setAuthor('Release Helpdesk')
SLIDES, SOURCES, QA = [], [], []
current = None


def rgb(value):
    return RGBColor.from_string(value)


def color(value):
    return tuple(int(value[i:i+2], 16) / 255 for i in (0, 2, 4))


def font_name(bold=False, mono=False):
    return 'Mono' if mono else 'NotoBold' if bold else 'Noto'


def segments(s, bold=False, mono=False):
    return [(v, 'Symbols' if v == '→' else font_name(bold, mono))
            for v in re.split('(→)', s) if v]


def text_width(s, size, bold=False, mono=False):
    return sum(pdfmetrics.stringWidth(v, f, size) for v, f in segments(s, bold, mono))


def rect(x, y, w, h, fill, stroke=None, radius=0):
    shape = current.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if radius:
        shape.adjustments[0] = min(radius / min(w, h), .5)
    if stroke:
        shape.line.color.rgb = rgb(stroke)
        shape.line.width = Pt(.7)
    else:
        shape.line.fill.background()
    pdf.setFillColorRGB(*color(fill))
    pdf.setStrokeColorRGB(*color(stroke or fill))
    pdf.setLineWidth(.7)
    pdf.roundRect(x, H-y-h, w, h, radius, stroke=bool(stroke), fill=1)
    return shape


def line(x1, y1, x2, y2, c=LINE, width=1.2, arrow=False):
    shape = current.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    shape.line.color.rgb, shape.line.width = rgb(c), Pt(width)
    pdf.setStrokeColorRGB(*color(c))
    pdf.setLineWidth(width)
    pdf.line(x1, H-y1, x2, H-y2)
    if arrow:
        angle = math.atan2(y2-y1, x2-x1)
        for delta in [-.48, .48]:
            xx, yy = x2-7*math.cos(angle+delta), y2-7*math.sin(angle+delta)
            line(x2, y2, xx, yy, c, width)


def wrap(s, maxw, size, bold=False, mono=False):
    font = font_name(bold, mono)
    result = []
    for para in s.split('\n'):
        if not para:
            result.append('')
            continue
        buf = ''
        for word in para.split():
            test = (buf + ' ' + word).strip()
            if text_width(test, size, bold, mono) <= maxw:
                buf = test
            else:
                if buf:
                    result.append(buf)
                buf = word
        result.append(buf)
    return result


def text(s, x, y, w, size=16, c=INK, bold=False, mono=False, leading=1.28, link=None):
    lines = wrap(s, w-3, size, bold, mono)
    height = len(lines)*size*leading+5
    box = current.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(height))
    tf = box.text_frame
    tf.clear()
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.word_wrap = False
    for i, value in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size*leading)
        for part, part_font in segments(value, bold, mono):
            r = p.add_run()
            r.text = part
            r.font.name = 'DejaVu Sans' if part_font == 'Symbols' else MONO if mono else FONT
            r.font.size, r.font.bold = Pt(size), bold
            r.font.color.rgb = rgb(c)
            if link:
                r.hyperlink.address = link
    pdf.setFillColorRGB(*color(c))
    pdf.setFont(font_name(bold, mono), size)
    for i, value in enumerate(lines):
        cursor = x
        for part, part_font in segments(value, bold, mono):
            pdf.setFont(part_font, size)
            pdf.drawString(cursor, H-y-size-i*size*leading, part)
            cursor += pdfmetrics.stringWidth(part, part_font, size)
    if link:
        pdf.linkURL(link, (x, H-y-height, x+w, H-y), relative=0, thickness=0)
    if y+height > 535 or x+w > 962:
        QA.append({'slide': len(SLIDES), 'text': s, 'bounds': [x,y,w,height]})
    return height


def pill(s, x, y, w, fill=PALE, c=BLUE):
    rect(x, y, w, 25, fill, radius=5)
    text(s, x+10, y+5, w-15, 10, c, True)


def image(name, x, y, w, h, label=None):
    path = ROOT / name
    im = Image.open(path)
    iw, ih = im.size
    scale = min(w/iw, h/ih)
    pw, ph = iw*scale, ih*scale
    px, py = x+(w-pw)/2, y+(h-ph)/2
    rect(x-1, y-1, w+2, h+2, WHITE, LINE, 3)
    pic = current.shapes.add_picture(str(path), Pt(px), Pt(py), width=Pt(pw), height=Pt(ph))
    pic.click_action.hyperlink.address = REPO+'/blob/main/'+name
    pdf.drawImage(str(path), px, H-py-ph, pw, ph, mask='auto')
    pdf.linkURL(REPO+'/blob/main/'+name, (x,H-y-h,x+w,H-y), relative=0, thickness=0)
    if label:
        text(label, x, y+h+7, w, 9, MUTED)
    SOURCES[-1].append(name)


def node(title, subtitle, x, y, w, h=65, fill=WHITE, accent=BLUE, dark=False):
    rect(x,y,w,h,fill, None if dark else LINE, 6)
    rect(x,y,3,h,accent)
    text(title,x+13,y+10,w-23,16,WHITE if dark else INK,True)
    if subtitle:
        text(subtitle,x+13,y+34,w-23,11,'B7CAE1' if dark else MUTED)


def point(num, title, body, x, y, w=300, dark=False):
    text(num,x,y,30,12,TEAL if not dark else '67DBCC',True)
    text(title,x+38,y-3,w-38,17,WHITE if dark else INK,True)
    text(body,x+38,y+25,w-38,13,'B7CAE1' if dark else MUTED)


def begin(section, title, sub='', dark=False, source=None):
    global current
    current = prs.slides.add_slide(prs.slide_layouts[6])
    SLIDES.append({'section': section, 'title': title, 'notes': ''})
    SOURCES.append([])
    rect(0,0,W,H,DARK if dark else PAPER)
    text('RH / RELEASE HELPDESK',46,22,350,9,'B7CAE1' if dark else MUTED,True)
    text(section.upper(),690,22,225,9,'B7CAE1' if dark else MUTED,True)
    rect(46,52,28,3,'67DBCC' if dark else BLUE)
    if title:
        text(title,46,69,870,31,WHITE if dark else INK,True,leading=1.13)
    if sub:
        text(sub,46,116,865,12,'B7CAE1' if dark else MUTED)
    line(46,507,914,507,'2C405A' if dark else LINE,.8)
    text('RELEASE ENGINEERING  /  PLATFORM WALKTHROUGH',46,518,360,8,'B7CAE1' if dark else MUTED)
    text(f'{len(SLIDES):02d} / 16',860,516,58,9,'B7CAE1' if dark else MUTED,True)
    if source:
        label, path = source
        url = path if path.startswith('https://') else REPO+'/blob/main/'+path
        text(label,46,489,820,8,'B7CAE1' if dark else MUTED,link=url)
        SOURCES[-1].append(path)


def end(notes, sources=()):
    SOURCES[-1].extend(sources)
    source_text = '\n'.join(s if s.startswith('https://') else REPO+'/blob/main/'+s for s in dict.fromkeys(SOURCES[-1]))
    current.notes_slide.notes_text_frame.text = notes+'\n\nEvidence / source links:\n'+source_text
    SLIDES[-1]['notes'] = notes
    pdf.showPage()


# 01 — a title with an actual product rather than a stock technology collage.
begin('Platform walkthrough', '', dark=True)
pill('BUILT · DEPLOYED · VERIFIED',46,85,217,'243F51','67DBCC')
text('Release\nHelpdesk',46,134,440,60,WHITE,True,leading=1.06)
text('From Git commit to\nmonitored Kubernetes.',48,282,420,23,'B7CAE1')
text('End-to-end release engineering',48,370,420,18,WHITE,True)
text('React · FastAPI · PostgreSQL · October 2026',48,401,410,11,'B7CAE1')
rect(48,443,258,31,BLUE,radius=5)
text('WATCH THE RECORDED DEMO',60,451,234,10,WHITE,True,link=DEMO_VIDEO)
image('images/aws-application-desktop.png',510,121,403,269)
text('ACTUAL AWS APPLICATION · CAPTURED 16 OCT 2026',510,403,403,8,'B7CAE1',True)
for x,big,small in [(510,'19','API test cases'),(650,'2 → 4 → 2','Ready replicas'),(812,'66','IaC resources')]:
    text(big,x,435,125 if x==510 else 150 if x==650 else 102,23,'67DBCC',True)
    text(small,x,465,140,10,'B7CAE1')
end('This project is Release Helpdesk, a service desk for deployment incidents, access requests and release follow-ups. I will show how a change moves from Git through tested and scanned container images into a monitored AWS EKS deployment. The AWS screenshots and measurements were captured on 16 October 2026. The cloud environment was then deliberately torn down after its evidence was preserved; the local stack remains reproducible.', ['README.md','docs/VERIFICATION.md',DEMO_VIDEO])

# 02 — the original domain and the complete user journey.
begin('01 / Application','A focused desk for release operations.','Capture the incident, assign an owner, and move the next release forward.',source=('Evidence · Real AWS browser walkthrough and responsive UI','images/README.md'))
image('images/aws-application-desktop.png',46,157,576,310,'Actual application through AWS ALB · full original screenshot')
point('01','Capture','Create a release ticket with context, priority and an assignee.',654,164,263)
point('02','Triage','Search and filter the queue; edit ownership and status.',654,263,263)
point('03','Resolve','Complete or delete tickets; see database-backed counts update.',654,362,263)
end('The original application domain is a release helpdesk. Tickets represent deployment incidents and operational requests. In the demo I create a ticket called Release rollback approval, assign Demo Student, select HIGH priority, and move it from TODO to IN_PROGRESS. I then filter the queue, show updated counts, edit and delete the ticket. The frontend has mobile, loading, empty and error states. Its activity feed is explicitly limited to the current browser session. The original TaskBoard API and field names are retained for compatibility. This is a classroom application using synthetic data, without end-user authentication. Show the supplied browser walkthrough or demonstrate the same flow locally.', ['images/aws-application-mobile.png','images/aws-browser-walkthrough.webm','application/frontend/src/main.jsx','docs/ATTRIBUTION.md'])

# 03 — native editable architectural diagram.
begin('02 / Architecture','One system. Three connected paths.','Request serving, release delivery and observability share the same deployed application.',source=('Source · Implemented architecture and routing','README.md'))
text('SERVE',46,161,100,10,BLUE,True)
node('Browser','Same origin',46,190,125)
node('AWS ALB','Ingress controller',211,190,160)
node('React / Nginx','2 frontend replicas',416,159,198)
node('FastAPI','2–4 backend replicas',416,254,198)
node('PostgreSQL 16','5 GiB gp3 PVC',711,254,203)
line(171,222,211,222,BLUE,1.8,True)
line(371,222,394,222,BLUE,1.8)
line(394,222,394,191,BLUE,1.8)
line(394,191,416,191,BLUE,1.8,True)
line(394,222,394,286,BLUE,1.8)
line(394,286,416,286,BLUE,1.8,True)
text('/',382,166,28,11,BLUE,True)
text('/api',376,263,45,10,BLUE,True)
line(614,286,711,286,BLUE,1.8,True)
text('SQLAlchemy',623,266,85,10,MUTED)
text('DELIVER',46,350,100,10,TEAL,True)
for x,title,desc in [(46,'GitHub','Commit SHA'),(223,'Actions','Test + build + scan'),(400,'GHCR','SHA-tagged images'),(577,'Helm + Flux','Reconciled release')]:
    node(title,desc,x,374,157,64,WHITE,TEAL)
    if x < 577:
        line(x+157,406,x+177,406,TEAL,1.6,True)
rect(772,365,142,81,'E3F2EF',radius=6)
text('OBSERVE',786,375,115,9,TEAL,True)
text('Prometheus\n→ Grafana',786,393,115,13,INK,True)
line(650,319,650,342,TEAL,1)
line(650,342,840,342,TEAL,1)
line(840,342,840,365,TEAL,1,True)
text('/metrics',720,324,90,9,TEAL)
end('Read the architecture as three paths. For requests, the AWS Load Balancer Controller implements the ALB Ingress: slash goes to React served by Nginx, and slash api goes to the FastAPI ClusterIP Service. Nginx also proxies API documentation to the backend. The API stores tickets in PostgreSQL on a persistent EBS volume. For delivery, GitHub Actions tests and scans source/dependencies/secrets, then builds and scans both images before publishing their full-SHA tags. Current GitOps promotion pins that verified SHA in Git and Flux reconciles the Helm release locally. The earlier cloud delivery used a short-lived AWS OIDC role and Helm. Terraform creates the VPC, EKS and IAM infrastructure before this application delivery flow. For observability, Prometheus scrapes the FastAPI metrics endpoint through a ServiceMonitor and Grafana visualizes those measurements. Arrows indicate logical flows; Terraform bootstrap and Alembic migration are explained on later slides.', ['helm/taskboard/templates/ingress.yaml','application/frontend/nginx.conf','.github/workflows/ci-cd.yml','monitoring/prometheus-values.yaml'])

# 04 — backend, test isolation and real persistence.
begin('03 / API, tests & data','Correctness starts before deployment.','A tested API contract and one migration authority protect the data layer.',source=('Evidence · PostgreSQL tests, migration and restart persistence','docs/VERIFICATION.md'))
rect(46,158,275,145,DARK,radius=7)
text('19',64,171,230,57,'67DBCC',True)
text('passing API test cases',65,243,230,16,WHITE,True)
text('Isolated test database',65,273,230,11,'B7CAE1')
text('SIX REST OPERATIONS',352,159,280,10,BLUE,True)
for i,s in enumerate(['GET    /api/tasks','GET    /api/tasks/{id}','GET    /api/tasks/stats','POST   /api/tasks','PUT    /api/tasks/{id}','DELETE /api/tasks/{id}']):
    text(s,352,188+i*27,330,12,INK,mono=True)
point('DB','Alembic only','Schema revision 0001_create_tasks; migrations run before traffic is ready.',700,166,217)
point('IO','Data survives','4 records survived a PostgreSQL Pod restart.',700,285,217)
rect(46,373,868,92,WHITE,LINE,6)
text('/health',63,389,160,16,BLUE,True,mono=True)
text('Process is alive',63,421,230,14,MUTED)
text('/ready',351,389,200,16,TEAL,True,mono=True)
text('Database + tasks table are usable',351,421,360,14,MUTED)
text('DB outage: 200 → 503 → 200',697,404,200,11,INK,True)
end('The backend implements six business REST operations, including list, detail, statistics and full create, update, delete behavior. Nineteen parametrized API cases cover CRUD, statistics, validation, missing records, documentation, metrics and database failure recovery. Unit tests default to isolated in-memory SQLite. CI and the integration runner use a separate PostgreSQL database that is migrated before tests run. The test suite does not silently create a PostgreSQL schema. Alembic is the deployment schema authority. Readiness checks the database and tasks table, while liveness checks the process. In an isolated real outage, ready changed from 200 to 503 and back to 200 while health stayed 200. Separately, the AWS PostgreSQL Pod restart changed the Pod UID while all four synthetic records remained on its PVC. Show Swagger CRUD and the saved PostgreSQL table screenshot while explaining this slide.', ['application/backend/tests/test_api.py','application/backend/tests/conftest.py','images/aws-postgresql-pytest-passed.png','images/aws-postgresql-table-and-persistence.png','evidence/logs/readiness-outage.txt','evidence/logs/aws-database-persistence.json'])

# 05 — reproducibility and deliberate image design.
begin('04 / Containers','The complete stack starts with one command.','Compose encodes dependencies; runtime containers execute as non-root users.',source=('Evidence · Real Compose stack, images and runtime users','images/local-compose-images-and-users.png'))
rect(46,158,868,48,DARK,radius=6)
text('docker compose up --build -d --wait',64,171,824,18,WHITE,mono=True)
for i,(a,b) in enumerate([('PostgreSQL','Healthy'),('Alembic','Completed'),('FastAPI','Ready'),('React / Nginx','Serving')]):
    xx=46+i*224
    node(a,b,xx,233,196,65)
    if i<3: line(xx+196,266,xx+224,266,BLUE,1.6,True)
rect(46,329,418,137,WHITE,LINE,7)
text('Backend runtime',64,345,370,18,INK,True)
text('Python 3.13 → Uvicorn',64,379,370,15,MUTED)
pill('USER 10001',64,420,119,'E3F2EF',TEAL)
rect(490,329,424,137,WHITE,LINE,7)
text('Frontend multi-stage build',508,345,385,18,INK,True)
text('Node 24 build → unprivileged Nginx',508,379,385,15,MUTED)
pill('USER 101',508,420,101,'E3F2EF',TEAL)
text('Only compiled assets enter the frontend runtime.',632,425,265,10,MUTED)
end('The full local stack runs using docker compose up with build and wait. Dependencies are explicit: PostgreSQL must be healthy, the migration must complete, the API must be ready, and then the frontend starts. The backend runs as UID 10001. The frontend uses a Node build stage and copies only static output into unprivileged Nginx running as UID 101 on port 8080. Local host ports are loopback-bound: the UI is on localhost 3000 and the API documentation is on localhost 8000 slash docs. PostgreSQL is not published to a host port. A named database volume keeps data across normal restarts. Show the saved real Compose image and user screenshot, then open the local application if demonstrating live.', ['docker-compose.yml','docker/backend.Dockerfile','docker/frontend.Dockerfile','evidence/logs/compose-persistence.txt'])

# 06 — the actual guarded pipeline and reviewed artifact promotion.
begin('05 / Continuous delivery','A guarded path from commit to cluster.','Tests and source-security gates must succeed before either image is published.',source=('Evidence · Actual integrated source-security and image publication run',SECURITY_RUN))
image('images/final/ci-security-release.png',46,159,487,276,'Actual hosted run · test, source security and publication passed')
for num,title,body,yy in [('01','Validate','PostgreSQL tests, frontend build and source/dependency/secret gates',162),('02','Build, scan & publish','Both images → Trivy → GHCR with immutable full-SHA tags',258),('03','Promote through Git','Pin the verified SHA → Flux → Helm → drift correction',354)]:
    point(num,title,body,567,yy,347)
text('Meaningful, traceable Git history',46,469,250,10,TEAL,True)
text('PRs validate without publishing; AWS deploy remains disabled after teardown.',320,469,590,10,TEAL,True)
end('The current GitHub Actions workflow validates tests and security independently, then requires both jobs before building and publishing images. The security job runs Bandit, Semgrep, pip-audit, npm audit and Gitleaks over full history plus working text. It writes fresh report artifacts for that execution. Tests use a clean migrated PostgreSQL service and also check the frontend, Compose, Helm, GitOps manifests, Terraform and the account guard. Both final images pass the Trivy HIGH/CRITICAL gate before the exact scanned images are pushed with the full source SHA. Pull requests cannot log in to the registry, publish or obtain AWS deployment credentials. The historical OIDC cloud release was verified separately. Current promotion changes the Git-owned HelmRelease to a tested image SHA, and Flux performs the rollout. The screenshot is run 39583131700 for image-source commit 1178b68; later evidence/documentation commits do not change the deployed image source.', ['.github/workflows/ci-cd.yml','evidence/logs/ci-security-release.json','evidence/logs/ci-final-source','evidence/logs/ci-final-images','gitops/releases/release.yaml'])

# 07 — source, dependencies, secrets and final-image gates.
begin('06 / Security','Security findings lead to repairs, then a release.','Source analysis, dependency audits and secret checks run before image publication.',dark=True,source=('Evidence · Fresh source and image reports for release 1178b68',SECURITY_RUN))
for x,label in [(46,'BACKEND IMAGE'),(490,'FRONTEND IMAGE')]:
    rect(x,164,424,166,DARK2,radius=7)
    text(label,x+20,181,380,11,'B7CAE1',True)
    text('0',x+20,209,90,65,'67DBCC',True)
    text('HIGH / CRITICAL\nfindings reported',x+128,228,270,18,WHITE,True)
text('Certifi trust-store advisory → upgrade to 2026.10.19 → rebuild and rescan',46,350,865,14,'B7CAE1')
line(46,385,914,385,'2C405A')
point('01','Source + dependencies','Bandit / Semgrep; pip-audit / npm audit.',46,408,430,True)
point('02','Secrets + final images','Full-history Gitleaks; both Trivy severity gates.',500,408,414,True)
end('The refreshed scans cover four layers. Bandit checks Python source and migrations; checked-in Semgrep rules check selected dangerous frontend patterns. pip-audit checks the complete pinned Python dependency set, while npm audit includes the frontend build/development graph. Gitleaks scans full reachable Git history and working text with redacted reports. Trivy checks OS and language packages in both final images and fails all HIGH/CRITICAL findings, including unfixed findings. The new dependency audit found the certifi trust-store advisory published after the earlier cloud run. The dependency was upgraded to 2026.10.19 and the images rebuilt. Both new images reported zero HIGH/CRITICAL findings at the recorded scan. Exact public CPython signing fingerprints in archived Trivy metadata have a narrow path/value exception; application secrets are not excluded. Synthetic unsafe Python, unsafe JavaScript and a fake secret each triggered a failed gate. These are dated results and do not claim immunity to future advisories or all application weaknesses.', ['security/README.md','security/gitleaks.toml','evidence/logs/security-final/summary.json','evidence/logs/security-final/negative-controls.json','evidence/logs/security-initial/pip-audit.json','evidence/logs/ci-final-images'])

# 08 — network topology with the actual classroom tradeoff.
begin('07 / Infrastructure as code','AWS infrastructure is declared in Terraform.','Mumbai / ap-south-1 · EKS 1.35 · two managed m7i-flex.large workers',source=('Evidence · Actual approved plan, apply and AWS Console captures','docs/VERIFICATION.md'))
rect(46,159,562,306,WHITE,LINE,7)
text('VPC 10.20.0.0/16 · TWO AVAILABILITY ZONES',62,173,520,11,BLUE,True)
for x,label,pub,priv in [(62,'ZONE A','10.20.101.0/24','10.20.1.0/24'),(340,'ZONE B','10.20.102.0/24','10.20.2.0/24')]:
    rect(x,205,250,238,PAPER,LINE,6)
    text(label,x+12,216,225,10,MUTED,True)
    rect(x+12,244,226,69,'E8EFFB',radius=5)
    text('Public subnet',x+24,252,200,13,BLUE,True)
    text(pub,x+24,276,200,11,MUTED,mono=True)
    rect(x+12,338,226,83,'E3F2EF',radius=5)
    text('Private subnet',x+24,349,200,13,TEAL,True)
    text(priv,x+24,372,200,11,MUTED,mono=True)
    text('Managed worker',x+24,395,200,10,TEAL)
text('ALB across public subnets · single NAT for egress',65,445,530,9,MUTED)
text('66',652,159,240,60,BLUE,True)
text('Terraform resources created',653,244,260,15,INK,True)
point('01','Pinned configuration','Modules, providers, add-ons and node AMI are versioned.',652,283,264)
point('02','Explicit account guard','Approved Free-plan credit use; no Paid-plan upgrade.',652,376,264)
end('Terraform provisioned the real AWS environment in ap-south-1. The network spans two availability zones with two public and two private subnets. The ALB uses the public subnets; managed workers use the private subnets. A single NAT gateway supports private egress as a classroom cost tradeoff, rather than claiming full zonal resilience. The EKS cluster used Kubernetes 1.35 and two eligible m7i-flex.large AL2023 workers. Terraform also configures IAM roles, OIDC, EBS CSI and cluster add-ons. The recorded apply created sixty-six Terraform resources. Versions and provider locks are retained. The approved run used account credits under the active Free plan; it did not upgrade the account. Show the saved actual Terraform plan and AWS VPC/EKS Console screenshots. Infrastructure is now removed, and recreation needs a newly approved plan and budget.', ['terraform/main.tf','terraform/variables.tf','terraform/.terraform.lock.hcl','terraform/terraform.tfvars.example','images/aws-terraform-plan.png','images/aws-vpc-console.png','images/aws-eks-active-console.png','images/aws-terraform-apply-and-ready-nodes.png'])

# 09 — deployment, storage, routing, migration discipline.
begin('08 / Kubernetes & Helm','Helm turns the stack into a repeatable release.','One versioned chart connects workloads, storage, probes, Services and Ingress.',source=('Evidence · Actual AWS Pods, Services, PVC, Helm release and Ingress','images/aws-kubernetes-helm-ingress.png'))
image('images/aws-kubernetes-helm-ingress.png',46,160,520,278,'Real AWS resource listing · taskboard namespace')
point('01','Replicated workloads','2 frontend Pods; backend HPA starts at 2 and can reach 4.',603,165,311)
point('02','Persistent database','PostgreSQL StatefulSet + encrypted 5 GiB gp3 volume.',603,268,311)
point('03','Controlled release','One migration Job with advisory lock; readiness gates traffic.',603,371,311)
text('Ingress: / → frontend:8080     /api → backend:8000',46,468,865,12,BLUE,True,mono=True)
end('The Helm chart deploys into the taskboard namespace. The frontend has two replicas, and the backend HPA maintains at least two replicas. Both Services use ClusterIP networking. The Ingress routes slash to frontend port 8080 and slash api to backend port 8000 through the actual ALB controller. PostgreSQL runs as a StatefulSet with a bound five-GiB encrypted gp3 volume provisioned through EBS CSI. A single migration hook Job runs the Alembic upgrade under a PostgreSQL advisory lock; it is not run concurrently in every API replica. On first install the hook runs after the database resources exist, and on upgrades it runs before the update. The deployment script explicitly checks StatefulSet and Deployment rollouts, then performs external readiness and CRUD checks. Explain Helm values, release history and rollback, noting that destructive database changes would require a compatible migration strategy.', ['helm/taskboard/Chart.yaml','helm/taskboard/values.yaml','helm/taskboard/values-prod.yaml','helm/taskboard/templates/migrate.yaml','application/backend/app/migrate.py','scripts/deploy.sh','evidence/logs/aws-persistent-storage.json'])

# GitOps — an actual controller, observed promotion and automatic repair.
begin('09 / GitOps','Committed intent keeps the running service honest.','Flux reconciles this application and Helm chart from the public repository on local Kubernetes.',source=('Evidence · Actual Flux promotion, drift recovery and digest checks','evidence/logs/gitops-final/summary.json'))
image('images/final/gitops-verification.png',46,159,520,278,'Actual controller state and recorded verification · local kind')
text('2 → 1 → 2',603,158,311,32,BLUE,True)
text('Manual frontend drift → automatic repair',605,204,309,12,MUTED)
point('01','Promote from Git','Scanned image SHA + gitops-v2 header arrived through Flux.',603,247,311)
point('02','Automatic recovery','2 frontend replicas restored in 34 seconds.',603,333,311)
text('Registry digests match both running application images.',46,469,865,12,TEAL,True)
end('Flux source-controller fetches this public repository. A Kustomization applies the Git-owned HelmRelease, and helm-controller installs or upgrades the same application chart. The baseline served gitops-v1. Commit 871c284 pinned both newly scanned 1178b68 image tags and changed the release marker to gitops-v2. The Nginx configuration checksum triggered a real frontend rollout, and HTTP exposed the new header. The verification then changed frontend replicas from two to one and observed Flux restore two in 34 seconds, without a forced reconcile, manual Helm upgrade or apply as the repair. The baseline exercise had independently recovered in 37 seconds. Backend replicas remain HPA-owned. The running image IDs were checked against public GHCR manifest digests and matched for both components. This GitOps verification is local kind; the earlier AWS environment remains deliberately removed.', ['gitops/README.md','gitops/releases/release.yaml','evidence/logs/gitops-baseline/summary.json','evidence/logs/gitops-final/summary.json','evidence/logs/gitops-final/image-identity.json'])

# 10 — end-to-end change, actual before and after.
begin('10 / End-to-end release','One small change proves the delivery chain.','A visible triage-help change reached the real AWS application through CI.',source=('Evidence · Successful AWS update for commit 179ad20 · click to open',RUN))
pill('BEFORE · 3ffe809',46,157,175)
pill('AFTER · 179ad20',505,157,175,'E3F2EF',TEAL)
image('images/aws-live-update-before.png',46,196,409,240)
image('images/aws-live-update-after.png',505,196,409,240)
line(464,311,495,311,TEAL,2.4,True)
text('Commit → tests → build → scan → SHA-tagged GHCR images → OIDC / Helm → visible update',46,455,868,14,INK,True)
end('This is the end-to-end release proof required by the presentation rubric. The before screenshot comes from the previous deployed application commit 3ffe809. Commit 179ad20 clarified the visible help text to: Assign an owner, set a priority, and move release tickets through resolution. The linked real GitHub Actions run passed all three jobs and deployed that full SHA. Registry and Kubernetes inspection verified the deployed tags and digests. The after screenshot shows the new text through the AWS ALB. During the recording, open the commit, the successful run and these two original images at readable zoom. These are preserved screens from the actual release; the AWS endpoint has since been torn down. A current push still runs CI, but deployment is intentionally disabled after cleanup.', ['images/aws-live-update-before.png','images/aws-live-update-after.png','evidence/logs/aws-registry-image-verification.json',REPO+'/commit/179ad20c5e4995bf458a33531905d6f0e694b580'])

# 11 — plot directly from recorded observations, not a decorative curve.
begin('11 / Autoscaling','CPU demand produced real scaling and recovery.','Observed on AWS: ready backend replicas increased under load, then recovered automatically.',source=('Evidence · Chart drawn from actual timestamped HPA observations','evidence/logs/aws-hpa-internal/frames.json'))
frames = json.loads((ROOT/'evidence/logs/aws-hpa-internal/frames.json').read_text())
summary = json.loads((ROOT/'evidence/logs/aws-hpa-internal/summary.json').read_text())
rect(46,159,580,280,WHITE,LINE,7)
text('READY BACKEND REPLICAS',65,175,500,10,BLUE,True)
cx, cy, cw, ch = 92, 217, 502, 174
for value in [2,3,4]:
    yy = cy+ch-(value-1)/3*ch
    line(cx,yy,cx+cw,yy,LINE,.8)
    text(str(value),68,yy-8,22,11,MUTED)
line(cx,cy,cx,cy+ch,MUTED,.8)
line(cx,cy+ch,cx+cw,cy+ch,MUTED,.8)
for seconds in [0,60,120,180,240,300,360]:
    xx=cx+seconds/360*cw
    text(str(seconds),xx-10,cy+ch+8,38,9,MUTED)
text('Elapsed seconds · sampled observations',191,418,400,9,MUTED)
coords=[]
for row in frames:
    value=row.get('ready_replicas')
    if value is not None:
        coords.append((cx+row['elapsed_seconds']/360*cw,cy+ch-(value-1)/3*ch))
for (x1,y1),(x2,y2) in zip(coords,coords[1:]):
    line(x1,y1,x2,y1,BLUE,2.5)
    line(x2,y1,x2,y2,BLUE,2.5)
for xx,yy in [coords[0],next(pt for pt,row in zip(coords,[r for r in frames if r.get('ready_replicas') is not None]) if row['ready_replicas']==4),coords[-1]]:
    rect(xx-3,yy-3,6,6,TEAL,radius=3)
text('2 → 4 → 2',665,159,249,35,BLUE,True)
text('Ready replicas, measured',667,207,247,12,MUTED)
for yy,big,label in [(248,'68,466','successful requests'),(312,'0','failed load requests'),(376,'180.1 s','24-worker in-cluster load')]:
    text(big,667,yy,247,27,INK,True)
    text(label,667,yy+37,247,11,MUTED)
text('HPA: 60% CPU target · 100m request / Pod · min 2, max 4 · peak 442% of requested CPU',46,459,868,11,TEAL,True)
end('This chart is derived from the actual saved AWS HPA observations. It plots ready backend replicas rather than merely desired replicas, so it proves usable capacity increased. A bounded in-cluster load ran twenty-four workers for 180.1 seconds, completing 68,466 requests with zero load-client failures. Ready replicas moved from two to four and later recovered automatically to two. The HPA target stayed at sixty percent of the existing 100-millicore CPU request, and observed CPU peaked at 442 percent of that request. That percentage is relative to the resource request, not percent of an entire worker node. An earlier external load completed without errors but did not generate enough backend CPU to scale; that honest non-scaling result is retained. The follow-up client ran in an isolated namespace, and its CPU was excluded from the backend HPA. These are observations of this specific test, not a general throughput or production-capacity guarantee.', ['evidence/logs/aws-hpa-internal/summary.json','evidence/logs/aws-hpa/summary.json','images/aws-hpa-before-load.png','images/aws-hpa-scaled-under-load.png','images/aws-hpa-recovered.png','helm/taskboard/templates/hpa.yaml'])

# 12 — telemetry from real traffic.
begin('12 / Observability','See traffic, latency and errors in one place.','Application instrumentation → ServiceMonitor → Prometheus → authenticated Grafana.',source=('Evidence · Authenticated AWS Grafana dashboard with populated panels','images/aws-grafana-live-dashboard.png'))
image('images/aws-grafana-live-dashboard.png',46,159,569,304,'Actual recorded AWS dashboard · HTTP traffic, errors, latency and replicas')
point('01','Instrument','Request counter + latency histogram; route-template labels.',651,164,263)
point('02','Scrape','Backend /metrics; real Prometheus application targets UP.',651,265,263)
point('03','Inspect privately','Grafana via private port-forward with generated credentials.',651,366,263)
end('The API exposes an actual HTTP request counter and a response-latency histogram at slash metrics. Handler labels use route templates, which avoids unbounded labels for every ticket ID. A ServiceMonitor selects the backend Service and named HTTP port, and Prometheus was verified scraping the application targets as UP. The screenshot is the real authenticated Grafana dashboard from the AWS run, with populated request rate, error rate, p95 latency and ready replica panels. Grafana stays a private ClusterIP Service and is accessed by authenticated port-forward. The small real error-rate sample includes the controlled PostgreSQL readiness interruption. Show the original metrics endpoint screenshot and Prometheus targets before opening this dashboard at full size.', ['application/backend/app/main.py','helm/taskboard/templates/servicemonitor.yaml','monitoring/dashboard.json','images/aws-metrics-endpoint.png','images/aws-prometheus-targets-up.png','evidence/logs/aws-prometheus-queries.json'])

# 13 — two diagnostic chains, not just failure screenshots.
begin('13 / Troubleshooting','Two failures. Two traceable recoveries.','Controlled exercises ran in an isolated namespace; diagnosis followed observable evidence.',source=('Evidence · Actual AWS failure symptoms, diagnosis and recovered HTTP responses','evidence/logs/aws-troubleshooting'))
for x,tag,title,chain,fix in [
    (46,'IMAGE FAILURE','ImagePullBackOff',['Pod not Ready','Describe Pod + events','Nonexistent image tag'],'Correct tag → rollout → Ready'),
    (490,'SERVICE FAILURE','No reachable backend',['Service has no endpoints','Compare selector + Pod labels','Selector does not match'],'Correct selector → HTTP success')]:
    rect(x,162,424,306,WHITE,LINE,7)
    text(tag,x+20,181,382,10,GOLD,True)
    text(title,x+20,209,382,23,INK,True)
    for i,s in enumerate(chain):
        text(f'{i+1:02d}',x+20,257+i*40,32,11,MUTED,True)
        text(s,x+57,254+i*40,342,15,INK)
        if i<2: line(x+32,276+i*40,x+32,287+i*40,LINE,1.4,True)
    rect(x+18,399,388,47,'E3F2EF',radius=5)
    text(fix,x+30,412,365,13,TEAL,True)
end('In the first isolated exercise, the Pod was not Ready because its image tag did not exist. Pod description and events exposed the registry pull failure. After correcting the tag, the rollout reached Ready. In the second, the Service selector did not match the Pod labels, so EndpointSlices contained no usable endpoints. Comparing the selector against labels identified the cause. Correcting the selector restored endpoints and a successful HTTP response. Both exercises ran in taskboard-troubleshooting and the recorded recovery is saved. In the video, open the original image and Service diagnosis screenshots and the final recovery screenshot, or show the saved terminal logs. Keep the explanation in the order symptom, evidence, cause, correction, verification.', ['troubleshooting/broken-image.yaml','troubleshooting/broken-service.yaml','scripts/troubleshooting.sh','images/aws-broken-image-diagnosis.png','images/aws-broken-service-diagnosis.png','images/aws-troubleshooting-recovered.png'])

# 14 — cleanup is a demonstrated engineering result.
begin('14 / Teardown & cost discipline','The deployment was verified, then removed.','Cloud evidence was pushed and checked before the authorized teardown.',source=('Evidence · Destroy completion and recorded residual-resource checks','evidence/logs/aws-cleanup-verification.json'))
image('images/aws-terraform-destroy-complete.png',46,161,510,270,'Actual Terraform destroy completion · 66 resources destroyed')
text('66',600,157,314,61,BLUE,True)
text('Terraform resources destroyed',602,235,312,17,INK,True)
point('01','Residual checks passed','Project EKS, VPC, ALB, workers, NAT and retained disks absent.',602,285,312)
point('02','Deployment disabled','DEPLOY_ENABLED=false; AWS application URL is historical.',602,380,312)
text('Recorded cleanup: 16 Oct 2026 · KMS key disabled / pending deletion · billing settlement may lag',46,466,868,10,MUTED)
end('The lifecycle ends with cleanup. All deployment screenshots and the walkthrough were pushed and verified remotely before deletion. The application and ALB were removed while the controller was still available; retained project storage was handled deliberately. Terraform then reported sixty-six resources destroyed. Read-only residual checks recorded the project EKS cluster, VPC, workers, ALB, target groups, NAT, EIP, network interfaces, retained disks and related roles as absent. The project KMS key is disabled and PendingDeletion for the required AWS waiting period, rather than already physically erased. Billing and credit reporting can lag, so the presentation does not claim zero usage or a settled cost total. DEPLOY_ENABLED is false and the old ALB URL is offline. The local stack remains available for review. No cloud resources need to be recreated to explain the preserved deployment evidence.', ['images/aws-cleanup-residual-checks.png','evidence/logs/aws-terraform-destroy.txt','evidence/logs/aws-evidence-before-teardown.json','docs/CLEANUP.md'])

# Outcome — engineering results and a practical production roadmap.
begin('15 / Outcomes','Tested. Reconciled. Observable.',dark=True,source=('Review · Executed results and reproducible project source','docs/VERIFICATION.md'))
for x,num,title,body in [
    (46,'01','Reproducible','Compose locally. Terraform for AWS. One chart for the application.'),
    (345,'02','Guarded','Source and dependency checks, secret scans and immutable image gates.'),
    (644,'03','Self-correcting','GitOps restores intent. HPA reacts to CPU. Metrics explain behavior.')]:
    rect(x,166,270,194,DARK2,radius=7)
    text(num,x+18,183,230,12,'67DBCC',True)
    text(title,x+18,217,240,20,WHITE,True)
    text(body,x+18,259,234,15,'B7CAE1')
text('Next production steps',46,389,800,17,WHITE,True)
text('End-user identity · HTTPS + DNS · managed database backups · service alerts',46,425,867,15,'B7CAE1')
text('github.com/NoteLegend/devOps-TaskBoard',46,466,820,13,'67DBCC',True,link=REPO)
end('The system connects product behavior to engineering evidence. The application is reproducible, data survives normal restarts, source and image gates stop unsafe changes, image SHAs connect source to registry artifacts, and Flux continuously reconciles the chosen release. The AWS lifecycle demonstrated real provisioning, ingress, persistent storage, monitoring, scaling and teardown. The later local GitOps verification adds real promotion and drift correction for the same application. Lessons are to verify application behavior after rollout, separate process liveness from database readiness, measure ready replicas for scaling, and tie every result to its revision and environment. Production improvements would add end-user authentication, HTTPS and DNS, managed database backups and service-level alerting. The repository includes original evidence, instructions and attribution; the recorded narration explains those results.', ['README.md','docs/VERIFICATION.md','docs/ATTRIBUTION.md'])

assert len(SLIDES)==16
assert not QA, f'Text bounds need review: {QA}'
prs.save(OUT/'Release-Helpdesk-Presentation.pptx')
pdf.save()

# Markdown companion keeps the script available even if an importer drops notes.
script = [
    '# Release Helpdesk — speaker notes and reference guide',
    '',
    '[Watch the recorded demo](' + DEMO_VIDEO + ') · [Presentation PDF](Release-Helpdesk-Presentation.pdf) · [PowerPoint](Release-Helpdesk-Presentation.pptx)',
    '',
    'Suggested length: **12–15 minutes**, including short screen demonstrations. The deck has 16 presentation slides. PowerPoint contains the same speaker notes.',
    '',
    '**Verified scope:** current source/security delivery and local Flux promotion/drift correction complement the preserved AWS lifecycle. Cloud screens are historical, and the cloud endpoint is offline. The recorded demo is available from the project overview; the remaining handover is submission of the published links.',
    '',
    '## Recording setup',
    '',
    '- Open `Release-Helpdesk-Presentation.pptx` in PowerPoint or LibreOffice Impress. It can also be imported into Google Slides with **File → Import slides → Upload**. A matching PDF is included for a consistent viewing copy.',
    '- Record your voice and screen at 1080p. Keep the slide or evidence window readable; pause briefly when switching windows.',
    '- Clearly say that the AWS deployment evidence was captured on **16 October 2026**, then the environment was removed after verification. Show the preserved AWS evidence and demonstrate the application locally. The old AWS URL is offline.',
    '- For a local app demo, run `docker compose up --build -d --wait` from the repository root. Open `http://localhost:3000` and `http://localhost:8000/docs`. Use synthetic tickets. Stop with `docker compose down` to retain the volume.',
    '- Open the public repository, successful AWS Actions run, GHCR package pages, and the `images/` folder before recording. Do not rerun Terraform, AWS load tests, or cloud troubleshooting for this presentation.',
    '- Hide credentials and unrelated private windows. The existing `images/aws-browser-walkthrough.webm` is a real browser recording, but it does not include your narration or cover the full capstone.',
    '',
    '## Suggested timing and screen changes',
    '',
    '| Segment | Slides | Time | Show during the recording |',
    '|---|---|---|---|',
    '| Introduction and product | 1–2 | 1:30 | Create/edit/filter/delete a local ticket; briefly show mobile evidence. |',
    '| Architecture, API and data | 3–4 | 1:45 | Swagger CRUD, statistics route and saved PostgreSQL table/persistence screenshot. |',
    '| Containers and CI/CD | 5–6 | 1:45 | Compose images/users, Git history, Actions jobs and both GHCR SHA tags. |',
    '| Security and infrastructure | 7–8 | 1:30 | Trivy reports, Terraform source/plan and saved AWS EKS/VPC screens. |',
    '| Kubernetes, GitOps and rollout | 9–11 | 2:00 | Helm resources, actual Flux promotion/drift correction, registry identity and historical AWS update. |',
    '| Scaling and monitoring | 12–13 | 1:45 | HPA before/scaled/recovered evidence, metrics, Prometheus UP and Grafana. |',
    '| Troubleshooting and cleanup | 14–15 | 1:30 | Both failure symptoms and recovery, then completed destroy. |',
    '| Outcomes | 16 | 0:30 | Summarize results, engineering lessons and the production roadmap. |',
    '',
    'The timings are flexible, not an assignment-imposed limit. Keep screen demonstrations short and explain what the evidence proves. The instructor decides whether preserved cloud evidence satisfies a live-demo request; this guide follows the assignment’s recorded-walkthrough option and the documented approved AWS shutdown.',
    '',
    '## Evidence to open at readable size',
    '',
]
for i,(meta,sources) in enumerate(zip(SLIDES,SOURCES),1):
    script.extend([f'### Slide {i:02d} — {meta["title"] or "Release Helpdesk"}', '',meta['notes'],'','Evidence:',''])
    for src in dict.fromkeys(sources):
        url=src if src.startswith('https://') else '../'+src
        script.append(f'- [{src}]({url})')
    script.append('')
script.extend([
    '## Final submission', '',
    'The recorded demo, PowerPoint, PDF and source repository are linked from the main README. Submit those published links through the requested channel; the supporting evidence remains in the repository.', '',
    '## Attribution', '',
    'Release Helpdesk adapts Nensi Ravaliya’s MIT-licensed TaskBoard starter at commit `a4e39b6b3a54dab1efd386c75ab86bf1d8a23454`. See [project attribution](../docs/ATTRIBUTION.md) and [LICENSE](../LICENSE). All application, AWS, CI, monitoring and terminal screenshots embedded here are genuine repository evidence. The architecture diagrams and HPA chart are presentation graphics; the HPA chart is generated directly from the saved observations.',
])
(OUT/'Demo-and-Speaker-Notes.md').write_text('\n'.join(script)+'\n')
index = {'slides':16,'format':'16:9 / editable text, diagrams and chart; original embedded screenshots',
         'cloud_evidence_date':'2026-10-16','prepared_date':'2026-10-21',
         'image_source_commit':'1178b684e26c00d06d05b66653d54dd5cebd191b',
         'gitops_promotion_commit':'871c2848d3dc6cb4076de33f6e7bb82ec23cdd0c',
         'slides_index':[{'number':i,'title':s['title'] or 'Release Helpdesk','sources':list(dict.fromkeys(src))} for i,(s,src) in enumerate(zip(SLIDES,SOURCES),1)]}
(OUT/'sources.json').write_text(json.dumps(index,indent=2)+'\n')
print(f'Created {len(SLIDES)} slides, editable PPTX, matching PDF, speaker notes and source index.')
